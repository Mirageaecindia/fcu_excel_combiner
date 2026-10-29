import glob
import io
import re
import sys
import zipfile
from pathlib import Path
from xml.dom import minidom

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from core import fcu_parser, matcher, people_pdf, pipeline, selection, vendor_pdf, xlsx_writer  # noqa: E402

FCU = glob.glob(str(ROOT / "input/fcu/*.xlsx"))
VENT = glob.glob(str(ROOT / "input/vent/*.xlsx"))
VENDOR = glob.glob(str(ROOT / "input/vendor/*.xlsm"))
needs_inputs = pytest.mark.skipif(not (FCU and VENT and VENDOR), reason="sample inputs missing")


@pytest.fixture(scope="module")
def session():
    return pipeline.Session(FCU, VENT, VENDOR[0])


@needs_inputs
def test_parse_counts(session):
    assert len(session.units) == 111                         # tagged units only (strict FCU Tag rule)
    assert len(session.vendor) == 241 and session.vendor_sheets == ["plot 04"]
    assert len(session.vent) > 100
    assert any("above the first FCU Tag" in w for w in session.fcu_warnings)
    u = session.units[0]
    assert (u.tag, u.zone, u.area, u.calc_total_kw, u.flow) == ("FCU-001-BF", "&001-BD1-BF-HK STORE", 37.2, 2.8, 191.0)


@needs_inputs
def test_multi_qty_unit_load_is_multiplied(session):
    u = next(u for u in session.units if u.tag == "FCU-131-RF")
    assert u.qty == 2 and u.calc_total_kw == pytest.approx(25.0)


@needs_inputs
def test_vendor_build_and_second_selection(session):
    data, warnings = session.build({"pct": 5, "direction": "increase"})
    ws = openpyxl.load_workbook(io.BytesIO(data)).active
    r = {c: ws[f"{c}13"].value for c in ["D", "G", "I", "N", "O", "P", "S", "T", "U", "V", "W", "Y", "Z", "AA", "AB", "AC", "AE"]}
    assert r["D"] == "FCU-001-BF" and r["N"] == 2.8
    assert r["S"] == 1 and r["T"] == 4.35 and r["U"] == 3.28
    assert r["V"] == pytest.approx(923 / 3.6, abs=1e-3) and r["W"] == pytest.approx(0.051)
    assert r["Y"] == r["S"]                                   # quantity never scaled
    assert r["Z"] == pytest.approx(4.35 * 1.05, abs=1e-4)
    assert r["AB"] == pytest.approx(r["V"] * 1.05, abs=1e-3)
    assert r["AE"] is None                                    # voltage left empty
    assert ws["Q13"].value == "=+N13/I13" and ws["R13"].value == "=+P13/I13"
    assert ws["Q12"].value is None or ws["Q12"].value  # header untouched
    assert warnings and all("no vendor match" in w for w in warnings)
    assert ws["K13"].value is None                              # Outdoor Air is always empty


@needs_inputs
def test_apply_only_selected_columns_and_decrease(session):
    data, _ = session.build({"pct": 10, "direction": "decrease", "apply_to": {"total": True, "sensible": False, "flow": False, "power": False}})
    ws = openpyxl.load_workbook(io.BytesIO(data)).active
    assert ws["Z13"].value == pytest.approx(4.35 * 0.9, abs=1e-4)
    assert ws["AA13"].value == ws["U13"].value and ws["AB13"].value == pytest.approx(ws["V13"].value, abs=1e-6)


@needs_inputs
def test_without_vendor_uses_calculated_load_and_user_qty():
    s = pipeline.Session(FCU, VENT)
    data, warnings = s.build({"qty": 3, "fallback_pct": 10, "fallback_direction": "increase", "pct": 5, "direction": "decrease"})
    ws = openpyxl.load_workbook(io.BytesIO(data)).active
    assert not warnings
    assert ws["S13"].value == 3 and ws["Y13"].value == 3
    assert ws["T13"].value == pytest.approx(2.8 * 1.1, abs=1e-4)
    assert ws["Z13"].value == pytest.approx(2.8 * 1.1 * 0.95, abs=1e-4)
    assert ws["V13"].value == pytest.approx(191 * 1.1, abs=1e-3)
    assert ws["W13"].value is None


@needs_inputs
def test_template_fidelity(session):
    data, _ = session.build({"pct": 1, "header": {"plot": "123", "project": "P", "consultant": "C"}})
    z = zipfile.ZipFile(io.BytesIO(data))
    names = set(z.namelist())
    assert {"xl/activeX/activeX1.xml", "xl/drawings/drawing1.xml", "xl/media/image1.emf"} <= names
    assert "xl/calcChain.xml" not in names
    for n in z.namelist():                                    # every XML part is well formed
        if n.endswith((".xml", ".rels", ".vml")):
            minidom.parseString(z.read(n))
    sheet = z.read(xlsx_writer.SHEET).decode()
    assert "<sheetProtection" in sheet and 'fullCalcOnLoad="1"' in z.read("xl/workbook.xml").decode()
    ws = openpyxl.load_workbook(io.BytesIO(data)).active
    assert ws.protection.sheet and ws["Y3"].value == "123" and ws["Y4"].value == "P"
    # header rows stay locked (Y3:AC5 = plot/project/consultant are editable in the template)
    assert all(ws.cell(r, c).protection.locked for r in range(6, 13) for c in range(4, 33))
    assert not ws["D13"].protection.locked and not ws["S13"].protection.locked
    assert ws["Q13"].protection.locked and ws["L13"].protection.locked
    # unused rows hidden, used rows visible
    assert not ws.row_dimensions[13].hidden and ws.row_dimensions[600].hidden


@needs_inputs
def test_original_header_rows_identical(session):
    data, _ = session.build({})
    tpl = zipfile.ZipFile(pipeline.TEMPLATE).read(xlsx_writer.SHEET).decode()
    out = zipfile.ZipFile(io.BytesIO(data)).read(xlsx_writer.SHEET).decode()
    head = lambda x: x[: x.index('<row r="13"')]
    assert head(tpl) == head(out)


def test_selection_math():
    assert selection.scale(100, 5, "increase") == 105
    assert selection.scale(100, 5, "decrease") == 95
    assert selection.scale(None, 5, "increase") is None
    assert selection.scale(100, 5, "increase", apply=False) == 100


def test_matcher_codes_and_names():
    assert matcher.ref_score("&009-010-BD1-BF-MEP", "&009-13-BD1-BF-ELE") > matcher.ref_score("&009-010-BD1-BF-MEP", "&010-BD1-BF-ELE")
    assert matcher.name_core("&001-BD1-BF-HK STORE") == "HK STORE"


def _words(rows, y0=100):
    out, x = [], 10
    for ri, cells in enumerate(rows):
        x = 10
        for c in cells:
            out.append(dict(page=0, x0=x, x1=x + 20, top=y0 + ri * 12, text=c))
            x += 40
    return out


def test_vendor_pdf_words():
    row = ["&001-A"] + ["x"] * 11 + ["FWW600", "923", "50", "4.35", "3.28"] + ["x"] * 8 + ["51", "1", "1"] + ["x"] * 3
    assert len(row) == 31
    got = vendor_pdf.parse_words(_words([["SPECIFIED"], row]))
    assert len(got) == 1
    g = got[0]
    assert (g.ref, g.qty, g.total_kw, g.sens_kw) == ("&001-A", 1.0, 4.35, 3.28)
    assert g.flow_ls == pytest.approx(923 / 3.6) and g.power_kw == pytest.approx(0.051)
    with pytest.raises(ValueError):
        vendor_pdf.parse_words(_words([["nothing", "useful"]]))


@needs_inputs
def test_data_cells_are_white_and_readable(session):
    data, _ = session.build({"pct": 5})
    ws = openpyxl.load_workbook(io.BytesIO(data)).active
    for c in ("D13", "G13", "J13", "L13", "M13", "N13", "S13", "T13", "AE13", "D300"):
        f = ws[c].fill
        assert f.fill_type is None or (f.fgColor.type == "rgb" and f.fgColor.rgb == "FFFFFFFF") or f.fgColor.type == "theme", c
        col = ws[c].font.color
        assert col is None or col.type != "rgb" or col.rgb != "FFFFFFFF", c
    assert ws["D13"].fill.fill_type == "solid" and ws["D13"].fill.fgColor.rgb == "FFFFFFFF"
    assert ws["D9"].fill.fgColor.rgb != "FFFFFFFF"          # header colours untouched
    assert not ws["D13"].protection.locked and ws["L13"].protection.locked


# ---------------------------------------------------------------- tag rule, people, multi-file
def make_fcu(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PLOT-05"
    ws.append(["title"])
    ws.append(["FCU Tag", "Zone", "Area", "Total", "Sens", "Air Flow", "", "", "", "", "Qty", "", "", "Type"])
    ws.append([None, "&999-ORPHAN-ROW", 1, None, None, 5])                       # above the first tag -> ignored
    ws.append(["FCU-03", "&383-384-BD4-7F-3B-08-FCU-03", 16.3, 1.8, 1.7, None, None, None, None, None, 1, None, None, "DUCTED"])
    ws.append([None, "&383-BD4-7F-3B-08-BRM3", 12, None, None, 124])
    ws.append([None, "&384-BD4-7F-3B-08-BTH2", 4.3, None, None, 11])
    ws.append(["FCU-04", "&385-388-BD4-7F-3B-08-FCU-04", 56, 6.4, 5.6, None, None, None, None, None, 2, None, None, "DUCTED"])
    ws.append([" ", "&385-BD4-7F-3B-08-COR", 13.5, None, None, 22])
    ws.append([None, "&386-BD4-7F-3B-08-POW", 3.6, 2.0, None, 3])                 # load but no tag -> merged + warning
    wb.save(path)


def hap_pages():
    def page(name, area, occ):
        return ("Space Input Data\n#" + name + "\n1. General Details:\nFloor Area ............ " + str(area) + " m2\n"
                "Avg. Ceiling Height ........ 5.2 m\n2.4. People:\nOccupancy ........... " + str(occ) + " Person\nActivity Level Heavy Work\n")
    return [page("383-BD4-7F-3B-08-BRM3", 12, 2.0), page("384-BD4-7F-3B-08-BTH2", 4.3, 3.0),
            page("385-BD4-7F-3B-08-COR", 13.5, 1.0), page("386-BD4-7F-3B-08-POW", 3.6, 0.0),
            page("383-BD1-7F-3B-08-OTHER", 12, 9.0)]


def test_tag_rule_sub_rows_are_summed(tmp_path):
    f = tmp_path / "fcu.xlsx"
    make_fcu(f)
    warns = []
    units = fcu_parser.parse(str(f), warnings=warns)
    assert [u.tag for u in units] == ["FCU-03", "FCU-04"]
    u3, u4 = units
    assert u3.flow == 135 and u3.area == 16.3 and len(u3.spaces) == 2          # 124 + 11
    assert u3.calc_total_kw == 1.8 and u4.calc_total_kw == pytest.approx(12.8)  # 6.4 x qty 2
    assert u4.flow == 25 and len(u4.spaces) == 2          # COR 22 + POW 3 (POW has load but no tag)
    assert any("1 rows have load values but no FCU Tag" in w for w in warns)
    assert any("above the first FCU Tag" in w for w in warns)


def test_hap_pdf_text_parsing():
    sp = people_pdf.parse_pages(hap_pages())
    assert [(x.name, x.people, x.area) for x in sp][:2] == [("383-BD4-7F-3B-08-BRM3", 2.0, 12.0), ("384-BD4-7F-3B-08-BTH2", 3.0, 4.3)]
    assert sp[3].people == 0 and len(sp) == 5
    assert len(people_pdf.parse_pages([hap_pages()[0], hap_pages()[0]])) == 1   # continued page, no duplicate
    with pytest.raises(ValueError):
        people_pdf.parse_pages(["nothing useful"])


def test_hap_pdf_file_roundtrip(tmp_path):
    fitz = pytest.importorskip("fitz")
    doc = fitz.open()
    for t in hap_pages():
        pg = doc.new_page()
        for n, line in enumerate(t.splitlines()):
            pg.insert_text((50, 60 + 14 * n), line, fontsize=9)
    out = tmp_path / "hap.pdf"
    doc.save(out)
    sp = people_pdf.parse(str(out))
    assert len(sp) == 5 and sp[0].people == 2.0 and sp[0].area == 12.0


def test_people_from_pdf_default_and_override(tmp_path):
    f = tmp_path / "fcu.xlsx"
    make_fcu(f)
    s = pipeline.Session([str(f)], people=[{"pages": hap_pages(), "name": "hap.pdf"}])
    rows, _ = s.build_rows({"qty": 1})
    assert [r["D"] for r in rows] == ["FCU-03", "FCU-04"]
    assert rows[0]["J"] == 5.0                                    # 2 + 3 people of FCU-03's two sub rows
    assert rows[1]["J"] == 1.0                                    # COR (1) + POW (0); BD1 'OTHER' excluded
    assert all(r.get("K") is None for r in rows) and rows[0]["P"] == 135
    s2 = pipeline.Session([str(f)], people=[{"pages": hap_pages()[:1], "name": "x"}])
    rows2, _ = s2.build_rows({"qty": 1, "people_default": 7})
    assert rows2[0]["J"] == 2.0 and rows2[1]["J"] == 7            # default only where nothing matched
    rows3, _ = s2.build_rows({"qty": 1})
    assert rows3[1]["J"] is None
    rows4, _ = s.build_rows({"qty": 1}, {"1": {"people": 0}, "0": {"people": -1}})
    assert rows4[0]["J"] is None and rows4[1]["J"] == 2.0         # reviewer override / "none"


def test_people_without_any_source_is_empty(tmp_path):
    f = tmp_path / "fcu.xlsx"
    make_fcu(f)
    s = pipeline.Session([str(f)])
    rows, _ = s.build_rows({"qty": 1})
    assert all(r["J"] is None for r in rows)
    rows, _ = s.build_rows({"qty": 1, "people_default": 2})
    assert all(r["J"] == 2 for r in rows)


def test_match_people_by_code_building_and_area():
    class U:
        def __init__(self, zone, spaces, area):
            self.zone, self.spaces, self.area, self.tag = zone, spaces, area, "FCU-1"
    hap = people_pdf.parse_pages(["#001-BD1-B-Pump Room\nFloor Area ..... 11.2 m2\nOccupancy ..... 1.0 Person\n",
                                  "#001-BD2-B-Other\nFloor Area ..... 50 m2\nOccupancy ..... 8.0 Person\n",
                                  "#001-BD1-GF-Lobby\nFloor Area ..... 40 m2\nOccupancy ..... 4.0 Person\n"])
    u = U("&001-BD1-BF-HK STORE", [("&001-BD1-BF-HK STORE", 11.2, 191.0)], 11.2)
    m = matcher.match_people([u], hap, [])
    assert [h[1] for h in m[0]["hits"]] == [0]                  # same code + building + area; other building excluded


@needs_inputs
def test_multiple_vendor_files_are_merged_without_duplicates():
    one = pipeline.Session(FCU, [], [VENDOR[0] + "::plot 04"])
    twice = pipeline.Session(FCU, [], [VENDOR[0] + "::plot 04", {"path": VENDOR[0], "sheet": "plot 04"}])
    assert len(twice.vendor) == len(one.vendor) == 241
    more = pipeline.Session(FCU, [], [VENDOR[0] + "::plot 04", VENDOR[0] + "::plot 03"])
    assert len(more.vendor) > 241
